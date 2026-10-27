"""평가 결과 보고용 엑셀 (인사담당자 요청 2026-10-22 — "바로 보고할 수 있을 정도로").

app.py 가 bp 를 붙인다. 세션 검사는 auth._guard 가 모든 /api/* 에 하므로 여기서는 HR 인지만 본다.

  GET  /api/eval/report?year=  실제 평가 (평가 3단계 ②): 서버가 그해 평가를 DB 에서 읽어 아래 엑셀을 만든다 (HR 만 · 공개 전에도 — 설계 9장)
                           거절 순서: HR 403 → 해(2026~2100, ASCII 숫자 네 자리) 400 → FIRST_YEAR 전 · 시작 안 한 해 404 → 제외 안 된 사람이 없음 409.
                           사람 줄은 evaluation._row_json 을 지난다(HR = full) — 칸 · 글자는 화면 evReportData 와 같고, clean() · build() 도 아래와 같다.
                           파일 이름 "{year}년 인사평가 결과.xlsx", no-store. 감사로그 EVAL_REPORT (해뿐 — 등급 · 이름 없음)
  POST /api/eval/report    화면이 보낸 평가 결과 → 같은 엑셀 (HR 만) — 시연 화면용 (가상 자료). 몸통의 ratio {"SA", "B", "D"}(0~70 정수)는 없으면 10 · 10 · 10

  엑셀 시트 ① 요약 — 숫자 카드 · 본부별 등급 분포 · 평가 묶음별 분포와 정원 기준 · 본부별 그래프 (이름 없음, A4 세로 한 장)
                  정원 기준 칸 = 그해 비율 · 늘린 정원이 들어간 숫자, 대표이사가 늘렸으면 뒤에 "(대표이사 허용 +n명 · 늘린 이유)" — 300자까지 받고 칸 폭보다 길면 줄을 바꾼다
  엑셀 시트 ② 명단 — 한 사람 한 줄 (대외비). 칸 제목 사번 · 이름 · 최종등급은 연봉 인상 계산(RZ_GRADE)이 그대로 읽는다

꾸민 표 · 그래프는 openpyxl 이 만든다(화면만으로는 그래프를 넣기 어렵다).
"""

from __future__ import annotations

import io
import re
import sqlite3
from urllib.parse import quote

from flask import Blueprint, Response, g, jsonify, request
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import auth
import evaluation
import kpi
from hr_eval.domain.models import Actor
from hr_eval.domain.quota import DEFAULT_RATIOS, QuotaMode, calculate_department_quota

bp = Blueprint("eval_report", __name__)
bp.register_error_handler(kpi._Refuse, kpi._refused)

MAX_BYTES, MAX_PEOPLE, MAX_POOLS, MAX_TEXT = 1_000_000, 3_000, 30, 80
MAX_QUOTA = 300   # 정원 기준 글자 — 늘린 이유(200자 — eval_rules.TEXT_MAX["back"])가 들어간다
GRADES, WAIT = ["S", "A", "B", "C", "D"], "처리 중"   # 처리 중 = 이의 처리 중이거나 아직 등급이 없음 (화면 evExportGrade 가 비워 보낸다)
TEXT_KEYS = ("div", "team", "no", "name", "title", "pool", "stage", "appeal")
DAY = re.compile(r"\d{4}-\d{2}-\d{2}")
CTRL = re.compile(r"[\x00-\x1f\x7f]")

FONT = "맑은 고딕"
ACCENT, INK, SUB, LINE, SOFT, CARD = "255DAA", "1B2433", "6B7685", "C9D2DE", "EAF2FB", "F4F7FB"   # 포털 색 (--accent)
GRADE_FILL = {"S": ("1F6B47", "FFFFFF"), "A": ("CDEBDD", "1F6B47"), "B": ("E3F4EC", "1F6B47"), "C": ("FCE9DF", "9A3A1F"), "D": ("F6C9B5", "9A3A1F")}
SERIES = {"S": "1F6B47", "A": "4CA57A", "B": "A7DBC0", "C": "F2B8A0", "D": "C8553A", WAIT: "BFC6D1"}
THIN = Side(style="thin", color=LINE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", indent=1)
ZERO_DASH = '0;-0;"-"'
EV_STAGES = {"SELF": "자기평가 중", "SCORE": "평가자 점수 대기", "GRADE": "본부장 등급 대기", "HR": "HR 확정 대기", "CEO": "대표이사 확인 대기", "DONE": "확정"}   # 화면 EV_STAGES (public/index.html) 와 같은 글자


# ----- 받은 값 검사 -----

def _text(v: object, most: int = MAX_TEXT) -> str | None:
    return v.strip() if isinstance(v, str) and len(v) <= most and not CTRL.search(v) else None


def _int(v: object, lo: int, hi: int) -> int | None:
    return v if isinstance(v, int) and not isinstance(v, bool) and lo <= v <= hi else None


def clean(body: object) -> dict | str:
    """화면이 보낸 것 → build() 가 쓰는 모양. 틀리면 이유(문자열)."""
    if not isinstance(body, dict):
        return "형식이 맞지 않습니다."
    year = _int(body.get("year"), 2000, 2100)
    day, released = body.get("date"), body.get("released")
    if year is None or not (isinstance(day, str) and DAY.fullmatch(day)) or not (released is None or (isinstance(released, str) and DAY.fullmatch(released))):
        return "연도 · 날짜가 맞지 않습니다."
    ratio = body.get("ratio", dict(zip(evaluation.RATIO_KEYS, DEFAULT_RATIOS)))   # 그해 등급 비율 — 요약 아래 설명 줄 (없으면 지금 규칙 10 · 10 · 10)
    ratio = {k: _int(ratio.get(k), 0, 70) for k in evaluation.RATIO_KEYS} if isinstance(ratio, dict) else None
    if ratio is None or None in ratio.values():
        return "등급 비율이 맞지 않습니다."
    appeals, open_ = _int(body.get("appeals"), 0, MAX_PEOPLE), _int(body.get("open"), 0, MAX_PEOPLE)
    people, pools = body.get("people"), body.get("pools")
    if appeals is None or open_ is None or not isinstance(people, list) or not isinstance(pools, list) or not 0 < len(people) <= MAX_PEOPLE or len(pools) > MAX_POOLS:
        return "인원 · 묶음이 맞지 않습니다."
    out = []
    for p in people:
        if not isinstance(p, dict):
            return "명단 형식이 맞지 않습니다."
        row = {k: _text(p.get(k, "")) for k in TEXT_KEYS}
        total, grade = p.get("total"), p.get("grade")
        if None in row.values() or not row["name"] or grade not in GRADES + [""]:
            return "명단에 맞지 않는 값이 있습니다."
        if total is not None and not (isinstance(total, (int, float)) and not isinstance(total, bool) and 0 <= total <= 200):
            return "총점이 맞지 않습니다."
        out.append({**row, "total": total, "grade": grade})
    pl = []
    for k in pools:
        name, quota = (_text(k.get("name")), _text(k.get("quota"), MAX_QUOTA)) if isinstance(k, dict) else (None, None)
        if not name or quota is None or not isinstance(k.get("ok"), bool):
            return "묶음에 맞지 않는 값이 있습니다."
        pl.append({"name": name, "quota": quota, "ok": k["ok"]})
    return {"year": year, "date": day, "released": released, "ratio": ratio, "appeals": appeals, "open": open_, "people": out, "pools": pl}


# ----- 엑셀 -----

def _f(size: int = 10, bold: bool = False, color: str = INK) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color)


def _fill(c: str) -> PatternFill:
    return PatternFill("solid", start_color=c, end_color=c)


def _title(ws, text: str, sub: str, last_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws["A1"], ws["A1"].font, ws["A1"].alignment = text, _f(18, True), Alignment(vertical="center")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws["A2"], ws["A2"].font = sub, _f(10, color=SUB)
    for c in range(1, last_col + 1):
        ws.cell(row=2, column=c).border = Border(bottom=Side(style="medium", color=ACCENT))
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 34, 20


def _header(ws, row: int, labels: list[str]) -> None:
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font, c.fill, c.alignment, c.border = _f(10, True, "FFFFFF"), _fill(ACCENT), CENTER, BOX
    ws.row_dimensions[row].height = 22


def _section(ws, row: int, text: str) -> None:
    ws.cell(row=row, column=1, value=text).font = _f(12, True, ACCENT)
    ws.row_dimensions[row].height = 22


def _counts(people: list[dict]) -> list[int]:
    return [sum(1 for p in people if (p["grade"] or WAIT) == g) for g in GRADES + [WAIT]]


def _print(ws, landscape: bool, one_page: bool) -> None:
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation, ws.page_setup.paperSize = ("landscape" if landscape else "portrait"), ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, (1 if one_page else 0)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.oddFooter.center.text = "&P / &N"


def _summary(ws, d: dict) -> None:
    ppl, last = d["people"], 10
    _title(ws, f"{d['year']}년 인사평가 결과 보고", f"기준일 {d['date']}   ·   결과 공개 {d['released'] or '공개 전'}   ·   인사포털", last)
    for i, w in enumerate([24, 8, 8, 8, 8, 8, 9, 9, 14, 11], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # 숫자 카드 (A · B:C · D:E · F:G · H:J)
    n, done, top = len(ppl), sum(1 for p in ppl if p["stage"] == "확정"), sum(1 for p in ppl if p["grade"] in ("S", "A"))
    cards = [("평가 대상", f"{n}명"), ("확정", f"{done}명"), ("S·A 비율", f"{top / n:.0%}"), ("이의 신청", f"{d['appeals']}건"), ("처리 중", f"{d['open']}건")]
    for (a, b), (lab, val) in zip([(1, 1), (2, 3), (4, 5), (6, 7), (8, 10)], cards):
        for r, v, font in ((4, lab, _f(9, color=SUB)), (5, val, _f(16, True, ACCENT))):
            if a != b:
                ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
            ws.cell(row=r, column=a, value=v).font = font
            ws.cell(row=r, column=a).alignment = CENTER
            for col in range(a, b + 1):
                ws.cell(row=r, column=col).fill = _fill(CARD)
                ws.cell(row=r, column=col).border = Border(left=THIN if col == a else None, right=THIN if col == b else None,
                                                           top=THIN if r == 4 else None, bottom=THIN if r == 5 else None)
    ws.row_dimensions[4].height, ws.row_dimensions[5].height = 20, 32

    # 1. 본부별 등급 분포 — 본부 순서는 받은 명단 순서 (화면이 조직 순서로 보낸다)
    _section(ws, 7, "1. 본부별 등급 분포")
    _header(ws, 8, ["본부", *GRADES, WAIT, "합계", "S·A 비율", "C·D 비율"])
    divs = list(dict.fromkeys(p["div"] for p in ppl))
    r = 9
    for dv in [*divs, "합계"]:
        sub = ppl if dv == "합계" else [p for p in ppl if p["div"] == dv]
        cnt = _counts(sub)
        for i, v in enumerate([dv or "(본부 없음)", *cnt, len(sub), (cnt[0] + cnt[1]) / len(sub), (cnt[3] + cnt[4]) / len(sub)], 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border, c.alignment, c.number_format = BOX, LEFT if i == 1 else CENTER, "0%" if i >= 9 else ZERO_DASH
            c.font = _f(10, dv == "합계" or i == 8)
            if dv == "합계":
                c.fill = _fill(SOFT)
        ws.row_dimensions[r].height = 20
        r += 1
    for i, v in enumerate(["비율", *[x / n for x in _counts(ppl)], 1, None, None], 1):
        c = ws.cell(row=r, column=i, value=v)
        c.border, c.alignment, c.number_format = BOX, LEFT if i == 1 else CENTER, '0.0%;-0.0%;"-"'
        c.font, c.fill = _f(10, True, ACCENT), _fill(SOFT)
    ws.row_dimensions[r].height = 20

    # 2. 평가 묶음별 분포 · 정원 기준 (정원 기준 · 충족은 화면이 규칙(evQuota · evPoolCheck)으로 셈해 보낸다)
    r += 2
    _section(ws, r, "2. 평가 묶음별 분포 · 정원 기준")
    _header(ws, r + 1, ["평가 묶음", *GRADES, WAIT, "인원", "정원 기준", "기준 충족"])
    r += 2
    for pool in d["pools"]:
        sub = [p for p in ppl if p["pool"] == pool["name"]]
        wrap = len(pool["quota"]) > 14   # 정원 기준 칸(폭 14)보다 긴 글자(늘린 정원 · 이유)는 줄을 바꿔 다 보이게 — 줄 높이는 엑셀이 맞춘다
        for i, v in enumerate([pool["name"], *_counts(sub), len(sub), pool["quota"], "충족" if pool["ok"] else "확인 필요"], 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border, c.alignment, c.number_format = BOX, LEFT if i == 1 else WRAP if i == 9 and wrap else CENTER, ZERO_DASH
            c.font = _f(10, i == 8, ("1F6B47" if pool["ok"] else "C8553A") if i == 10 else INK)
        if not wrap:
            ws.row_dimensions[r].height = 20
        r += 1

    # 3. 그래프 — 숨은 칸(L~Q)에서 그린다: 0명은 비워 둬야 막대 위에 '-'가 안 나온다 (표의 서식이 글자로 따라온다)
    r += 1
    _section(ws, r, "3. 본부별 등급 분포 그래프")
    first, end, h = 9, 8 + len(divs), 12
    for i, g in enumerate([*GRADES, WAIT]):
        ws.cell(row=8, column=h + i, value=g)
        for rr in range(first, end + 1):
            ws.cell(row=rr, column=h + i, value=ws.cell(row=rr, column=2 + i).value or None)
        ws.column_dimensions[get_column_letter(h + i)].hidden = True
    ch = BarChart()
    ch.type, ch.grouping, ch.overlap, ch.gapWidth = "bar", "stacked", 100, 60
    ch.height, ch.width = max(5.0, 1.2 * len(divs) + 2), 23
    ch.legend.position, ch.visible_cells_only = "b", False
    ch.y_axis.majorGridlines, ch.y_axis.delete = None, True
    ch.x_axis.scaling.orientation, ch.x_axis.delete = "maxMin", False   # 표와 같은 순서 (위에서 아래로)
    ch.add_data(Reference(ws, min_col=h, max_col=h + 5, min_row=8, max_row=end), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=first, max_row=end))
    for s, g in zip(ch.series, [*GRADES, WAIT]):
        s.graphicalProperties.solidFill, s.graphicalProperties.line.solidFill = SERIES[g], "FFFFFF"
        cp = CharacterProperties(sz=900, b=True, solidFill="FFFFFF" if g in ("S", "D") else INK)
        s.dLbls = DataLabelList(showVal=True, showLegendKey=False, showCatName=False, showSerName=False, showPercent=False, showLeaderLines=False)
        s.dLbls.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])
    ws.add_chart(ch, f"A{r + 1}")

    note = r + 3 + int(ch.height / 0.5)   # 그래프 높이(cm) ≈ 줄 수 × 0.5
    ws.cell(row=note, column=1, value="· 처리 중 = 이의 신청을 처리하고 있거나 아직 등급이 정해지지 않은 사람입니다.").font = _f(9, color=SUB)
    rt = d["ratio"]
    ws.cell(row=note + 1, column=1, value=f"· 정원 기준 = 묶음 인원의 S·A {rt['SA']}% · B {rt['B']}% · D {rt['D']}%(반올림), 4명 이하 묶음은 S·A·B 합계 1명."
                                          " 대표이사가 늘린 정원은 묶음 줄에 적습니다.").font = _f(9, color=SUB)
    ws.print_options.horizontalCentered = True
    _print(ws, landscape=False, one_page=True)


def _roster(ws, d: dict) -> None:
    cols = [("No", 5), ("본부", 18), ("부서", 15), ("사번", 9), ("이름", 10), ("직급", 9), ("평가 묶음", 19),
            ("총점", 8), ("최종등급", 9), ("상태", 12), ("이의 신청", 20)]
    _title(ws, f"{d['year']}년 인사평가 결과 명단", f"기준일 {d['date']}   ·   {len(d['people'])}명   ·   대외비", len(cols))
    for i, (_, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    _header(ws, 4, [c for c, _ in cols])
    for k, p in enumerate(d["people"], 1):
        r = 4 + k
        vals = [k, p["div"], p["team"], p["no"] or None, p["name"], p["title"], p["pool"],
                "현장직" if p["total"] is None else p["total"], p["grade"] or None, p["stage"], p["appeal"] or None]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border, c.font, c.alignment = BOX, _f(10), CENTER if i in (1, 4, 6, 8, 9, 10) else LEFT
            if k % 2 == 0:
                c.fill = _fill("FAFBFD")
        ws.cell(row=r, column=8).number_format = "0.0"
        if p["total"] is None:   # 현장직은 점수 없이 팀장 평가
            ws.cell(row=r, column=8).font = _f(9, color=SUB)
        if p["grade"]:
            bg, fg = GRADE_FILL[p["grade"]]
            ws.cell(row=r, column=9).fill, ws.cell(row=r, column=9).font = _fill(bg), _f(10, True, fg)
        if p["appeal"]:
            ws.cell(row=r, column=11).font = _f(10, True, "C8553A")
        ws.row_dimensions[r].height = 19
    ws.auto_filter.ref = f"A4:{get_column_letter(len(cols))}{4 + len(d['people'])}"
    ws.freeze_panes = "A5"
    ws.print_title_rows = "4:4"
    _print(ws, landscape=True, one_page=False)


def build(d: dict) -> bytes:
    wb = Workbook()
    wb.active.title = "요약"   # 그래프가 시트 이름으로 칸을 가리키므로 먼저 정한다
    _summary(wb.active, d)
    _roster(wb.create_sheet("명단"), d)
    for ws in wb.worksheets:   # 받은 글자가 '=' 로 시작해도 수식이 아니라 글자로 둔다 (이 파일에 수식은 없다)
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f":
                    c.data_type = "s"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _file(d: dict, data: bytes) -> Response:
    """엑셀 응답 — 파일 이름은 한글 (filename* UTF-8)."""
    resp = Response(data, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp.headers["Content-Disposition"] = f"attachment; filename=eval_report_{d['year']}.xlsx; filename*=UTF-8''{quote(f"{d['year']}년 인사평가 결과.xlsx")}"
    return resp


@bp.post("/api/eval/report")
def report() -> Response:
    if (denied := auth._hr_only()) is not None:
        return denied
    if (request.content_length or 0) > MAX_BYTES:
        return jsonify(error="보낸 내용이 너무 큽니다."), 413
    d = clean(request.get_json(silent=True))
    if isinstance(d, str):
        return jsonify(error=d), 400
    return _file(d, build(d))


# ----- 실제 평가: DB 에서 -----

def _appeal_text(a: dict | None, grade: str | None) -> str:
    """이의 신청 칸 글자 (화면 evReportData): 없음 "" · 처리 중 (지금 등급) · 대표이사 확인 대기 (전 → 후 — 내려받는 HR 본인의 이의는 후가 null 이라 화살표 없이) ·
    등급이 바뀌었으면 "전 → 후 반영" · 그대로면 유지."""
    if a is None:
        return ""
    if a["status"] == "OPEN":
        return f"처리 중 (지금 {grade or ''})"
    if a["status"] == "CEO":
        return f"대표이사 확인 대기 ({a['before']} → {a['after']})" if a["after"] else "대표이사 확인 대기"
    return f"{a['before']} → {a['after']} 반영" if a["before"] != a["after"] else "유지"


def _from_db(conn: sqlite3.Connection, year: int, me: Actor) -> dict:
    """그해 평가를 DB 에서 읽어 clean() 이 받는 모양으로 — 칸 · 글자는 화면 evReportData 와 같다. 한 트랜잭션(읽는 사이에 단계가 바뀌지 않게).
    사람 줄은 모두 evaluation._row_json 을 지난다 (공개 전 등급을 자르는 곳이 한 곳이게 — 여기는 HR 만 오므로 full 이라 등급이 든다)."""
    with kpi._txn(conn):
        ctx = evaluation._ctx(conn, year)
        if not (ctx.round and ctx.round["started"]):
            raise kpi._Refuse("평가를 시작한 해가 아닙니다.", 404)
        rows = [r for r in sorted(ctx.rows.values(), key=lambda r: evaluation._order(ctx, r)) if r["out_at"] is None]   # 본부 → 팀 → 팀장 먼저 → 이름
        shown = [j for r in rows if (j := evaluation._row_json(ctx, me, r)) is not None]
        if not shown:
            raise kpi._Refuse("보고서에 넣을 분이 없습니다.", 409)
        no = dict(conn.execute("SELECT id, employee_no FROM users"))
        pools = []
        for k in sorted(ctx.pools, key=evaluation._pool_rank):   # 정원 글자 = 그해 비율 · 늘린 정원이 들어간 숫자 (늘렸으면 뒤에 "(대표이사 허용 +n명 · 늘린 이유)" — 이유의 줄바꿈 · 겹친 빈칸은 한 칸)
            _, errors, _, _, members = evaluation._pool(ctx, k)
            extra = evaluation._extra_of(ctx, k)
            q = calculate_department_quota(len(members), evaluation._ratios_of(ctx), extra)
            quota = f"S·A·B 합계 {q.upper}명" if q.mode is QuotaMode.GROUPED else f"S·A {q.a}명 · B {q.b}명 · D {q.d}명"
            why = " ".join((ctx.pools[k]["extra_reason"] or "").split())   # 늘렸으면 이유가 있다 (DB 트리거)
            pools.append({"name": evaluation._pool_label(ctx, k), "quota": quota + (f" (대표이사 허용 +{sum(extra)}명 · {why})" if any(extra) else ""), "ok": not errors})
    people = [{"div": j["div"], "team": j["team"] or evaluation.NO_TEAM, "no": no[j["pid"]] or "", "name": j["name"], "title": j["title"], "pool": j["poolName"],
               "total": None if j["total"] is None else round(j["total"], 1),
               "grade": "" if j["appeal"] and j["appeal"]["status"] != "DONE" else j["grade"] or "",   # 이의 처리 중이면 비운다 (화면 evExportGrade)
               "stage": EV_STAGES[j["stage"]].replace("등급", "배정"),   # '등급' 글자가 첫 사람 줄에 오면 연봉 인상 계산이 그 줄까지 제목으로 읽는다
               "appeal": _appeal_text(j["appeal"], j["grade"])} for j in shown]
    rel = ctx.round["released"]
    return {"year": year, "date": auth._today().isoformat(), "released": evaluation._day(rel).isoformat() if rel else None, "ratio": ctx.round["ratio"],
            "appeals": sum(1 for j in shown if j["appeal"]), "open": sum(1 for j in shown if j["appeal"] and j["appeal"]["status"] != "DONE"),
            "people": people, "pools": pools}


@bp.get("/api/eval/report")
def report_live() -> Response:
    if (denied := auth._hr_only()) is not None:
        return denied
    year = evaluation._asked(request.args.get("year", ""))
    if year < evaluation.FIRST_YEAR:   # 포털 밖에서 평가한 해
        raise kpi._Refuse("평가를 시작한 해가 아닙니다.", 404)
    conn = auth._db()
    d = clean(_from_db(conn, year, g.actor))
    if isinstance(d, str):   # 서버가 만든 자료라 일어나지 않아야 한다 — 그래도 500 대신 이유를 돌려준다
        raise kpi._Refuse(d, 409)
    data = build(d)
    evaluation._log(conn, "EVAL_REPORT", "ev_round", year, year=year)   # 만든 뒤에 남긴다 — 해뿐 (등급 · 이름 없음)
    resp = _file(d, data)
    resp.headers["Cache-Control"] = "no-store"
    return resp


if __name__ == "__main__":   # 견본: python eval_report.py 받은값.json 저장할파일.xlsx
    import json
    import sys
    v = clean(json.load(open(sys.argv[1], encoding="utf-8")))
    assert isinstance(v, dict), v
    open(sys.argv[2], "wb").write(build(v))
    print("saved", sys.argv[2])
